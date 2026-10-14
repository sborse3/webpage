#!/usr/bin/env python3
"""Update the original résumé template while preserving its visual formatting.

Requires python-docx, npm install, and the uploaded source DOCX in the repo root.
"""
from copy import deepcopy
from pathlib import Path
import json
import subprocess

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.text.paragraph import Paragraph
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[1]
DATA = json.loads(subprocess.check_output(['node', str(ROOT / 'scripts/export_resume_data.mjs')]))
TEMPLATE = ROOT / 'shubhankar_borse_2026_short.docx'
OUT = ROOT / 'public' / 'Shubhankar_Borse_Resume.docx'
GRAY, RED, GOLD, BLUE = '595959', 'C00000', 'C98B2E', '002060'
GUTTER = 1.25
ROLE_INDENT = GUTTER + .20
doc = Document(TEMPLATE)
prototypes = [deepcopy(p._p) for p in doc.paragraphs]

# Preserve the original centered name, photograph, QR code, rules, contact icons,
# theme, fonts, numbering definitions, margins and document relationships.
for p in list(doc.paragraphs)[10:]:
    p._p.getparent().remove(p._p)
doc.core_properties.title = 'Shubhankar Borse | Staff AI Research Scientist'
doc.core_properties.author = 'Shubhankar Borse'
doc.core_properties.subject = 'Generative and multimodal AI research and commercialization'


def run(p, value, bold=False, italic=False, color=GRAY, size=11):
    r = p.add_run(value)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor.from_string(color)
    r.bold, r.italic = bold, italic
    return r


def paragraph(proto=11, before=0, after=2, keep=False):
    el = deepcopy(prototypes[proto])
    # Keep only the original paragraph formatting, not stale fields or text.
    for child in list(el):
        if child.tag != qn('w:pPr'):
            el.remove(child)
    # Unique Word editing IDs are unnecessary for new paragraphs.
    for attr in list(el.attrib):
        del el.attrib[attr]
    doc._element.body.insert(len(doc._element.body) - 1, el)
    p = Paragraph(el, doc._body)
    f = p.paragraph_format
    f.left_indent, f.first_line_indent = Inches(GUTTER), Inches(0)
    f.space_before, f.space_after = Pt(before), Pt(after)
    f.keep_with_next, f.keep_together = keep, True
    f.widow_control = True
    return p


def section(label, first_text='', before=12):
    p = paragraph(38, before=before, after=5, keep=True)
    p.paragraph_format.left_indent = Inches(GUTTER)
    p.paragraph_format.first_line_indent = Inches(-GUTTER)
    p.paragraph_format.tab_stops.clear_all()
    p.paragraph_format.tab_stops.add_tab_stop(Inches(GUTTER))
    run(p, label.upper(), bold=True, color=RED)
    if first_text:
        run(p, '\t')
        run(p, first_text, bold=True)
    return p


def body(value, bold=False, italic=False, before=0, after=2, keep=False):
    p = paragraph(11, before, after, keep)
    run(p, value, bold, italic)
    return p


def bullet(value, label=None, publication=False, indent=GUTTER):
    p = paragraph(39 if publication else 12, after=4)
    p.paragraph_format.left_indent = Inches(indent + .25)
    p.paragraph_format.first_line_indent = Inches(-.25)
    p.paragraph_format.line_spacing = 1.15 if publication else 1.0
    if label:
        run(p, label, bold=True)
    run(p, value)
    return p


def hyperlink(p, title, url, color=GRAY, underline=False, size=11):
    h = OxmlElement('w:hyperlink')
    h.set(qn('r:id'), p.part.relate_to(url, RT.HYPERLINK, is_external=True))
    r, pr = OxmlElement('w:r'), OxmlElement('w:rPr')
    font = OxmlElement('w:rFonts')
    for key in ['ascii', 'hAnsi', 'cs']: font.set(qn('w:'+key), 'Times New Roman')
    pr.append(font)
    c = OxmlElement('w:color'); c.set(qn('w:val'), color); pr.append(c)
    sz = OxmlElement('w:sz'); sz.set(qn('w:val'), str(int(size * 2))); pr.append(sz)
    if underline:
        u = OxmlElement('w:u'); u.set(qn('w:val'), 'single'); pr.append(u)
    r.append(pr)
    t = OxmlElement('w:t'); t.text = title; r.append(t)
    h.append(r); p._p.append(h)


# Replace the summary in place, retaining its original italic, tracked typography.
summary = doc.paragraphs[3]
summary_style = deepcopy(summary.runs[0]._r.rPr)
summary.clear()
r = summary.add_run(
    'Generative AI, Multimodal Models, Video Generation, and Model Efficiency. '
    f"{DATA['portfolio']['papers']} papers, {DATA['portfolio']['filings']} patent filings, "
    f"{DATA['portfolio']['grants']} granted patents, and {DATA['portfolio']['citations']} citations."
)
r._r.insert(0, summary_style)

# Keep the original contact icons, but place the second row's icons inline so
# changing its labels cannot leave floating icons on top of the text.
contact = doc.paragraphs[7]
drawings = list(contact._p.iter(qn('w:drawing')))
icons = {}
for drawing in drawings:
    anchor = drawing.find(qn('wp:anchor'))
    description = anchor.find(qn('wp:docPr')).get('descr', '')
    inline = OxmlElement('wp:inline')
    for tag in ['wp:extent', 'wp:effectExtent', 'wp:docPr', 'wp:cNvGraphicFramePr', 'a:graphic']:
        child = anchor.find(qn(tag))
        if child is not None: inline.append(deepcopy(child))
    replacement = OxmlElement('w:drawing'); replacement.append(inline)
    icons['patents' if 'Head with gears' in description else 'webpage'] = replacement
contact.clear()
for i, (key, label, url) in enumerate([
    ('patents', 'Patents', DATA['portfolio']['patentUrl']),
    ('webpage', 'Webpage', DATA['contact']['website']),
]):
    if i: run(contact, '  /  ', size=10)
    contact.add_run()._r.append(icons[key])
    run(contact, ' ', size=10)
    hyperlink(contact, label, url, BLUE, True, size=10)
for rel_id, rel in list(doc.part.rels.items()):
    if rel.is_external and 'paperswithcode.com' in rel.target_ref:
        doc.part.drop_rel(rel_id)

# The original LinkedIn label looked like a link but had no target.
for p in doc.paragraphs[:10]:
    for r in list(p.runs):
        if r.text == DATA['contact']['linkedin']:
            h = OxmlElement('w:hyperlink')
            h.set(qn('r:id'), p.part.relate_to('https://linkedin.com/in/'+r.text, RT.HYPERLINK, is_external=True))
            r._r.addprevious(h); h.append(r._r)

periods = [job['period'] for job in DATA['experience']]
roles = ['Staff AI Research Scientist', 'Sr. Deep Learning Research Engineer',
         'Deep Learning Research Engineer', 'Machine Learning Algorithms Intern']


def outline(p, level):
    el = OxmlElement('w:outlineLvl')
    el.set(qn('w:val'), str(level))
    p._p.get_or_add_pPr().append(el)


# Employers sit at the body's left edge; roles and their bullets nest underneath.
p = section('Experience', before=0)
run(p, '\t')
run(p, 'Qualcomm AI Research, San Diego  ·  ', bold=True)
run(p, 'Feb 2019 – Present', italic=True, color=GOLD)
outline(p, 0)
p.paragraph_format.space_after = Pt(8)

for i, job in enumerate(DATA['experience'][:3]):
    p = paragraph(18, before=13 if i else 2, after=4, keep=True)
    p.paragraph_format.left_indent = Inches(ROLE_INDENT)
    outline(p, 1)
    run(p, roles[i]+' ', bold=True)
    run(p, periods[i], bold=True, italic=True, color=GOLD)
    scope = body(job['scope'], italic=True, after=5, keep=True)
    scope.paragraph_format.left_indent = Inches(ROLE_INDENT)
    for achievement in job['achievements']:
        bullet(achievement, indent=ROLE_INDENT)

intern = DATA['experience'][3]
p = paragraph(18, before=20, after=8, keep=True)
run(p, intern['company']+', '+intern['location']+'  ·  ', bold=True)
run(p, periods[3], italic=True, color=GOLD)
outline(p, 0)
p = paragraph(18, after=4, keep=True)
p.paragraph_format.left_indent = Inches(ROLE_INDENT)
run(p, roles[3], bold=True)
for achievement in intern['achievements']:
    bullet(achievement, indent=ROLE_INDENT)

first_authored = [p for p in DATA['publications'] if p['category'] in ('First Author', 'Co-First Author')]
selected_ids = ['hiphy', 'madi', 'zero-shot-peft', 'lora-x']
publications = {p['id']: p for p in DATA['publications']}


def paper(pub):
    p = bullet('', publication=True)
    hyperlink(p, pub['title'], pub['scholarLink'], color=BLUE, underline=True)
    venue = pub['venue'].replace(' · Datasets & Benchmarks', ' Datasets & Benchmarks')
    if pub.get('status') == 'Preprint': venue = 'arXiv preprint'
    suffix = venue+' '+pub['year']
    if pub['category'] == 'Co-First Author': suffix += '; co-first author'
    run(p, ' ('+suffix+')', bold=True)


section('Publications', 'First- & Co-First-Authored Papers')
for pub in first_authored: paper(pub)
section('Publications', 'Selected Co-Authored Papers')
for key in selected_ids: paper(publications[key])
p = body('', after=3)
hyperlink(p, 'Complete publication list on Google Scholar', DATA['contact']['scholar'], BLUE, True)

# Education follows both publication groups, using the source document's layout.
for i, edu in enumerate(DATA['education']):
    if i == 0:
        p = section('Education')
        run(p, '\t')
    else:
        p = paragraph(85, before=7, keep=True)
    degree = 'Master of Science in Electrical and Computer Engineering' if i == 0 else 'Bachelor of Technology in Electronics and Telecommunication Engineering'
    run(p, degree+' – ', bold=True)
    run(p, 'December 2018' if i == 0 else 'May 2017', bold=True, italic=True, color=GOLD)
    body(edu['institution'], keep=True)
    bullet(edu['gpa'], label='GPA ')
    for detail in edu['details']: bullet(detail)

section('Patents', f"{DATA['portfolio']['filings']} filings · {DATA['portfolio']['grants']} granted patents")
body('Research Areas:', bold=True, keep=True)
for label, value in [
    ('Generative AI: ', 'Diffusion models, zero-shot personalization, content–style composition, image editing, and blended output generation.'),
    ('Efficient Models & Adaptation: ', 'Fourier and sparse adapters, training-free adapter transfer, diffusion distillation, hardware-aware architectures, and test-time adaptation.'),
    ('Computer Vision: ', 'Panoptic and semantic segmentation, BEV perception, depth completion, point-cloud processing, and boundary-aware learning.'),
    ('Multimodal & Sensor Fusion: ', 'Camera–radar–LiDAR integration, occlusion handling, intermodal training, and adaptive confidence mechanisms.'),
    ('Security & Watermarking: ', 'Video and image watermarking, and adversarial robustness.'),
]: bullet(value, label)

section('Awards', 'Awards & Recognition')
for award in DATA['recognition']:
    title = award['title']
    if title == 'Top Reviewer': title = 'NeurIPS Top Reviewer'
    if title == 'Best Demo Award — Honorable Mention': title = 'CVPR Best Demo Award — Honorable Mention'
    if title == 'Qualstar Award': title = 'Qualcomm Qualstar Award'
    p = bullet('')
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    run(p, title+' ', bold=True)
    run(p, award['year'], bold=True, italic=True, color=GOLD)
    # Keep the informative award details; avoid repeating the award name in prose.
    if award['title'] not in ('Top Reviewer', 'Qualcomm CTO Award'):
        run(p, '\n'+award['description'])

section('Service', 'Reviewing & Academic Service')
for entry in DATA['reviewing']:
    label = entry['organization']
    if entry['period']: label += ' '+entry['period']
    if entry['organization'] == 'Multiple journals':
        bullet('Peer review for multiple journals.')
    else:
        bullet(entry['role']+'.', label=label+': ')
for activity in DATA['academicActivities']: bullet(activity)

OUT.parent.mkdir(exist_ok=True)
doc.save(OUT)
print(f'{OUT}\n{len(first_authored)} first/co-first-author papers; {len(selected_ids)} selected co-authored papers.')
